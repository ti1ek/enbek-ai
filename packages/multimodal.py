"""Multimodal document extraction for attachments.

Pulls text from a user-attached file so it can be fed to the LLM as context:
  - images (png/jpg/webp): vision-OCR via the chat model (genuine multimodality)
  - PDF / DOCX: LlamaParse (preserves complex tables/layout as markdown, OCRs
    scanned pages). Falls back to local PyMuPDF / python-docx when no
    LLAMA_CLOUD_API_KEY is set or the cloud call fails — and a scanned PDF in
    the fallback path is rendered to images and vision-OCR'd.
  - TXT / MD: plain text.

Vision calls go through packages.llm.chat_complete (OpenAI-primary /
Gemini-fallback). Both gpt-4.1(-mini) and gemini-2.5-flash are multimodal, so
the fallback path also handles images.
"""
from __future__ import annotations

import base64
import io

from packages.config import settings
from packages.llm import chat_complete

MAX_OUTPUT_CHARS = 12_000          # cap injected context to control tokens
MAX_PDF_VISION_PAGES = 8           # cap scanned-PDF pages sent to the vision model
_PDF_TEXT_MIN_CHARS = 40           # below this per doc → treat as scanned, use vision

_IMAGE_EXT = {"png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff"}
_IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "webp": "image/webp", "gif": "image/gif", "bmp": "image/bmp", "tiff": "image/tiff",
}
_LLAMAPARSE_EXT = {"pdf", "docx", "doc", "pptx", "xlsx"}

_OCR_PROMPT = (
    "Это изображение документа по трудовому праву Казахстана "
    "(трудовой договор, приказ, справка, уведомление и т.п.). "
    "Извлеки ВЕСЬ текст дословно, сохраняя номера пунктов, статей, даты, суммы, ФИО и реквизиты. "
    "Таблицы оформи в виде markdown-таблиц. "
    "Не комментируй и не анализируй — верни только извлечённый текст. "
    "Если текста на изображении нет — верни пустую строку."
)


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _vision_ocr(image_bytes: bytes, mime: str) -> str:
    """OCR a single image via the multimodal chat model."""
    b64 = base64.b64encode(image_bytes).decode()
    resp = chat_complete(
        model=settings.llm_mini_model,
        temperature=0.0,
        max_tokens=2000,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": _OCR_PROMPT},
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                ],
            }
        ],
    )
    return (resp.choices[0].message.content or "").strip()


def _llamaparse(data: bytes, filename: str) -> str:
    """Parse PDF/DOCX via LlamaParse → markdown (keeps tables, OCRs scans)."""
    from llama_cloud_services import LlamaParse

    parser = LlamaParse(
        api_key=settings.llama_cloud_api_key,
        result_type="markdown",
        language="ru",
        verbose=False,
    )
    docs = parser.load_data(data, extra_info={"file_name": filename})
    return "\n\n".join((d.text or "").strip() for d in docs).strip()


def _extract_pdf_local(data: bytes) -> tuple[str, str, int]:
    """Local PDF fallback: PyMuPDF text layer; scanned pages → vision-OCR."""
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    pages = doc.page_count
    text_layer = "\n".join(page.get_text() for page in doc).strip()

    if len(text_layer) >= _PDF_TEXT_MIN_CHARS:
        doc.close()
        return text_layer, "pdf_text", pages

    parts: list[str] = []
    for i, page in enumerate(doc):
        if i >= MAX_PDF_VISION_PAGES:
            parts.append(f"[…пропущено {pages - MAX_PDF_VISION_PAGES} страниц…]")
            break
        pix = page.get_pixmap(dpi=150)
        ocr = _vision_ocr(pix.tobytes("png"), "image/png")
        if ocr:
            parts.append(f"[стр. {i + 1}]\n{ocr}")
    doc.close()
    return "\n\n".join(parts).strip(), "pdf_vision_ocr", pages


def _extract_docx_local(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip()


def _extract_doc_like(data: bytes, filename: str, ext: str) -> tuple[str, str, int]:
    """PDF/DOCX/etc: LlamaParse first (best for tables), local fallback on miss/error."""
    if settings.llama_cloud_api_key:
        try:
            text = _llamaparse(data, filename)
            if text:
                return text, "llamaparse", 0
        except Exception:
            pass  # fall through to local extraction
    if ext == "pdf":
        return _extract_pdf_local(data)
    if ext == "docx":
        return _extract_docx_local(data), "docx", 0
    raise ValueError(
        f"Для .{ext} нужен LLAMA_CLOUD_API_KEY (локальный парсер недоступен)"
    )


def extract_document(filename: str, data: bytes) -> dict:
    """Extract text from an attached file.

    Returns {"filename", "text", "method", "pages", "chars", "truncated"}.
    Raises ValueError for unsupported types.
    """
    ext = _ext(filename)
    pages = 0

    if ext in _IMAGE_EXT:
        text = _vision_ocr(data, _IMAGE_MIME.get(ext, "image/png"))
        method = "vision_ocr"
    elif ext in _LLAMAPARSE_EXT:
        text, method, pages = _extract_doc_like(data, filename, ext)
    elif ext in {"txt", "md"}:
        text = data.decode("utf-8", errors="replace").strip()
        method = "text"
    else:
        raise ValueError(f"Неподдерживаемый тип файла: .{ext}")

    truncated = len(text) > MAX_OUTPUT_CHARS
    if truncated:
        text = text[:MAX_OUTPUT_CHARS] + "\n[…документ обрезан…]"

    return {
        "filename": filename,
        "text": text,
        "method": method,
        "pages": pages,
        "chars": len(text),
        "truncated": truncated,
    }
